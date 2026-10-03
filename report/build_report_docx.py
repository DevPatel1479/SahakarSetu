import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout_box(doc, text, title="KEY ARCHITECTURAL HIGHLIGHT"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border highlight
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="1E3A5F"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"★ {title}\n")
    run_t.bold = True
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    
    run_b = p.add_run(text)
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(4)

def build_docx():
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SahakarSetu — Technical Architecture Report | PS 26087")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Ministry of Cooperation | National Council for Cooperative Training (NCCT)")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    
    # Set base font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    # Title Page / Banner
    p_banner = doc.add_paragraph()
    p_banner.paragraph_format.space_before = Pt(10)
    p_banner.paragraph_format.space_after = Pt(4)
    r_sub = p_banner.add_run("SMART INDIA HACKATHON 2026 | PROBLEM STATEMENT ID: 26087\nMINISTRY OF COOPERATION | NATIONAL COUNCIL FOR COOPERATIVE TRAINING")
    r_sub.font.size = Pt(9.5)
    r_sub.bold = True
    r_sub.font.color.rgb = RGBColor(0xC1, 0x7F, 0x24)
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("SahakarSetu: AI & LMS-Enabled Cooperative Capacity Building, ERP & Employment Ecosystem")
    r_title.font.size = Pt(22)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(16)
    r_desc = p_desc.add_run("Comprehensive Technical Architecture & Solution Engineering Report\nHardware Category | Theme: Smart Education | Mode: Software + Hardware")
    r_desc.font.size = Pt(12)
    r_desc.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    
    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(16)
    r_div = p_div.add_run("―" * 52)
    r_div.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    
    # Read Markdown source
    md_path = "D:/SahakarSetu/docs/TECHNICAL_ARCHITECTURE.md"
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()
    
    # We will process section by section
    lines = md_text.splitlines()
    i = 0
    in_code = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        
        # Check code block
        if line.startswith("```"):
            if in_code:
                in_code = False
                # Output code box
                code_text = "\n".join(code_lines)
                code_tbl = doc.add_table(rows=1, cols=1)
                code_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                code_tbl.autofit = False
                code_tbl.columns[0].width = Inches(6.5)
                c_cell = code_tbl.cell(0, 0)
                set_cell_background(c_cell, "F8FAFC")
                set_cell_margins(c_cell, top=100, bottom=100, left=150, right=150)
                cp = c_cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                c_run = cp.add_run(code_text)
                c_run.font.name = "Consolas"
                c_run.font.size = Pt(8.5)
                c_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(2)
                p_spacer.paragraph_format.space_after = Pt(4)
                code_lines = []
            else:
                in_code = True
                code_lines = []
            i += 1
            continue
            
        if in_code:
            code_lines.append(line)
            i += 1
            continue
            
        # Table parsing
        if line.startswith("|") and i + 1 < len(lines) and lines[i+1].startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
                
            # Parse table rows
            headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
            # Skip separator line (line 1)
            data_rows = []
            for r in table_lines[2:]:
                cols = [c.strip() for c in r.split("|")[1:-1]]
                data_rows.append(cols)
                
            tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = True
            
            # Format header
            for col_idx, h_text in enumerate(headers):
                cell = tbl.cell(0, col_idx)
                set_cell_background(cell, "1E3A5F")
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(re.sub(r'\*\*(.*?)\*\*', r'\1', h_text))
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
            # Format data rows
            for row_idx, r_data in enumerate(data_rows):
                bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                for col_idx, cell_text in enumerate(r_data):
                    if col_idx < len(headers):
                        cell = tbl.cell(row_idx + 1, col_idx)
                        set_cell_background(cell, bg)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        # Clean bold markdown
                        cleaned = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text)
                        r = p.add_run(cleaned)
                        r.font.size = Pt(8.5)
                        if "**" in cell_text or col_idx == 0:
                            r.bold = True
                            
            p_spacer = doc.add_paragraph()
            p_spacer.paragraph_format.space_before = Pt(4)
            p_spacer.paragraph_format.space_after = Pt(6)
            continue
            
        # Skip top markdown header since we did banner
        if line.startswith("# ") and i < 5:
            i += 1
            continue
        if line.startswith("**Comprehensive Technical") or line.startswith("**Smart India") or line.startswith("**Problem Statement") or line.startswith("**Organization") or line.startswith("**Department") or line.startswith("**Theme:"):
            i += 1
            continue
            
        # Headings
        if line.startswith("## "):
            h_text = line.replace("## ", "").strip()
            h = doc.add_heading(h_text, level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            for r in h.runs:
                r.font.size = Pt(14)
                r.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
                
            # Embed diagram after Section 4
            if "End-to-End System Architecture" in h_text:
                img_path = "D:/SahakarSetu/architecture/system-architecture.png"
                if os.path.exists(img_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(4)
                    run_img = p_img.add_run()
                    run_img.add_picture(img_path, width=Inches(6.2))
                    
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_after = Pt(10)
                    r_cap = p_cap.add_run("Figure 1: SahakarSetu End-to-End Hybrid System Architecture (Cloud + Edge Box)")
                    r_cap.font.size = Pt(8.5)
                    r_cap.italic = True
                    r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            i += 1
            continue
            
        if line.startswith("### "):
            h_text = line.replace("### ", "").strip()
            h = doc.add_heading(h_text, level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            for r in h.runs:
                r.font.size = Pt(11.5)
                r.bold = True
                r.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
            i += 1
            continue
            
        if line.strip() == "---":
            i += 1
            continue
            
        # Bullet list
        if line.startswith("- ") or line.startswith("* "):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            # Parse bold spans
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue
            
        # Regular paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*|\$.*?\$)', line.strip())
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                elif part.startswith("$") and part.endswith("$"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Cambria Math"
                    r.italic = True
                else:
                    p.add_run(part)
                    
        i += 1
        
    out_docx = "D:/SahakarSetu/report/SahakarSetu_Technical_Architecture_Report.docx"
    doc.save(out_docx)
    print(f"Successfully generated DOCX report at: {out_docx}")

if __name__ == "__main__":
    build_docx()
